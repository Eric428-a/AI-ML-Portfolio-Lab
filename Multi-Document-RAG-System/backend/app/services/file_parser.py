import csv
import io
import logging
from pathlib import Path

from docx import Document as WordDocument
from fastapi import HTTPException, UploadFile
from pypdf import PdfReader

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".csv",
}

TEXT_EXTENSIONS = {".txt", ".md", ".csv"}


def validate_file_extension(filename: str) -> str:
    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=415,
            detail=(
                f"Unsupported file type '{extension}'. "
                f"Supported types: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
            ),
        )

    return extension


async def read_upload(upload: UploadFile, max_size_bytes: int) -> bytes:
    content = await upload.read(max_size_bytes + 1)

    if not content:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")

    if len(content) > max_size_bytes:
        raise HTTPException(
            status_code=413,
            detail=f"File exceeds the maximum size of {max_size_bytes // (1024 * 1024)} MB.",
        )

    return content


def extract_pdf(data: bytes) -> tuple[str, list[dict]]:
    try:
        reader = PdfReader(io.BytesIO(data), strict=False)

        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise ValueError("Encrypted PDF requires a password.")
            except Exception as exc:
                raise ValueError("Password-protected PDFs are not supported.") from exc

        page_texts = []
        pages = []

        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            page_texts.append(text)
            pages.append(
                {
                    "page": page_number,
                    "text": text,
                }
            )

        return "\n\n".join(page_texts), pages

    except Exception as exc:
        logger.exception("PDF extraction failed")
        raise ValueError(f"Could not read PDF: {exc}") from exc


def extract_docx(data: bytes) -> tuple[str, list[dict]]:
    try:
        document = WordDocument(io.BytesIO(data))
        sections = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                sections.append(paragraph.text.strip())

        for table_index, table in enumerate(document.tables, start=1):
            rows = []
            for row in table.rows:
                rows.append(" | ".join(cell.text.strip() for cell in row.cells))
            sections.append(
                f"Table {table_index}\n" + "\n".join(rows)
            )

        text = "\n\n".join(sections)
        return text, [{"section": "document", "text": text}]

    except Exception as exc:
        logger.exception("DOCX extraction failed")
        raise ValueError(f"Could not read DOCX: {exc}") from exc


def extract_text_file(data: bytes) -> tuple[str, list[dict]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("utf-8", errors="replace")

    return text, [{"section": "document", "text": text}]


def extract_csv(data: bytes) -> tuple[str, list[dict]]:
    try:
        decoded = data.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(decoded))
        rows = []

        for row_number, row in enumerate(reader, start=1):
            if any(cell.strip() for cell in row):
                rows.append(
                    f"Row {row_number}: " + " | ".join(cell.strip() for cell in row)
                )

        text = "\n".join(rows)
        return text, [{"section": "csv", "text": text}]

    except Exception as exc:
        logger.exception("CSV extraction failed")
        raise ValueError(f"Could not read CSV: {exc}") from exc


def extract_document_text(
    filename: str,
    data: bytes,
) -> tuple[str, list[dict]]:
    extension = validate_file_extension(filename)

    if extension == ".pdf":
        result = extract_pdf(data)
    elif extension == ".docx":
        result = extract_docx(data)
    elif extension == ".csv":
        result = extract_csv(data)
    else:
        result = extract_text_file(data)

    text, segments = result

    if not text.strip():
        raise ValueError(
            "No readable text was found. Scanned PDFs require OCR before ingestion."
        )

    return text, segments